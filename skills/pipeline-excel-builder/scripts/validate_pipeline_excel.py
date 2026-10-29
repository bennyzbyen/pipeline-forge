#!/usr/bin/env python3
"""Validate a DataHub Pipeline Export Excel workbook."""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import zipfile
from collections import defaultdict
from xml.etree import ElementTree as ET
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


DATA_UTILIZATION_SHEET = "Data Utilization"
TARGET_SHEET = "Target & Catalog"
FIELD_SHEET = "Target Field"
PIPELINE_SHEET = "Pipeline"

EXPECTED_HEADERS = {
    DATA_UTILIZATION_SHEET: ["*data_utilization_name", "data_utilization_description"],
    TARGET_SHEET: [
        "*data_utilization_name",
        "*target_name",
        "target_description",
        "data_storage_type",
        "db",
        "table_name",
        "table_description",
        "dataset_name",
        "dataset_title",
        "dataset_description",
        "dataset_business_owner",
        "dataset_business_owner_email",
        "dataset_it_owner",
        "dataset_it_owner_email",
        "dataset_fe",
        "dataset_fe_email",
        "dataset_it_bp",
        "dataset_it_bp_email",
        "dataset_data_engineer",
        "dataset_data_engineer_email",
        "dataset_source",
    ],
    FIELD_SHEET: [
        "*target_name",
        "*field_name",
        "*field_label",
        "field_description",
        "*field_type",
        "field_length",
        "field_sequence",
    ],
    PIPELINE_SHEET: [
        "*data_utilization_name",
        "*pipeline_name",
        "pipeline_description",
        "*enable",
        "*is_octopus",
        "pipeline_trigger",
        "pipeline_trigger_start",
        "pipeline_trigger_end",
        "pipeline_status_notification",
        "pipeline_notification_emails",
        "task1_name",
        "task1_description",
        "task1_link_target_names",
        "task1_mlp_params",
    ],
}

ALLOWED_STORAGE_TYPES = {"", "HBASE", "HDFS", "CLICKHOUSE", "SV_CLICKHOUSE", "MSSQL", "MYSQL"}
REQUIRED_FIELD_TYPE = "TEXT"
REQUIRED_FIELD_LENGTH = "200"
SHARED_STRINGS_PATH = "xl/sharedStrings.xml"


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def project_prefix(value: str) -> str:
    value = clean(value).lower()
    return value.split("_", 1)[0] if "_" in value else value


def is_project_owned_table(table_name: str, data_utilization_name: str) -> bool:
    prefix = project_prefix(data_utilization_name)
    return bool(prefix and clean(table_name).lower().startswith(prefix + "_"))


def headers(ws, count: int) -> list[str]:
    return ["" if ws.cell(2, col).value is None else str(ws.cell(2, col).value) for col in range(1, count + 1)]


def row_dicts(ws, expected_headers: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for row_idx in range(3, ws.max_row + 1):
        row = {header: clean(ws.cell(row_idx, col).value) for col, header in enumerate(expected_headers, start=1)}
        if any(row.values()):
            row["_row"] = str(row_idx)
            rows.append(row)
    return rows


def required_headers(expected_headers: list[str]) -> list[str]:
    return [header for header in expected_headers if header.startswith("*")]


def add_duplicate_errors(values: list[str], label: str, errors: list[str]) -> None:
    seen: set[str] = set()
    duplicates: set[str] = set()
    for value in values:
        if not value:
            continue
        if value in seen:
            duplicates.add(value)
        seen.add(value)
    for value in sorted(duplicates):
        errors.append(f"Duplicate {label}: {value}")


def add_empty_inline_string_errors(ws, max_col: int, errors: list[str]) -> None:
    for (row_idx, col_idx), cell in ws._cells.items():
        if row_idx < 3 or col_idx > max_col:
            continue
        if cell.value is None and cell.data_type == "inlineStr":
            errors.append(f"{ws.title} {cell.coordinate} is an empty inline string cell; leave blank cells absent/null.")


def add_ooxml_storage_errors(path: Path, errors: list[str]) -> None:
    with zipfile.ZipFile(path, "r") as zf:
        names = set(zf.namelist())
        if SHARED_STRINGS_PATH not in names:
            errors.append("Workbook is missing xl/sharedStrings.xml; the preserved export profile requires shared-string text cells.")
        inline_string_cells: list[str] = []
        for name in sorted(names):
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            xml = zf.read(name).decode("utf-8")
            count = xml.count('t="inlineStr"')
            if count:
                inline_string_cells.append(f"{name}: {count}")
        if inline_string_cells:
            errors.append(
                "Workbook contains inlineStr cells; preserve the export's sharedStrings profile (POI STRING itself also supports inline text). "
                + "; ".join(inline_string_cells)
            )


def unreadable_workbook_result(path: Path, exc: Exception) -> dict[str, Any]:
    detail = str(exc).strip() or exc.__class__.__name__
    return {
        "ok": False,
        "workbook": str(path),
        "row_counts": {},
        "errors": [f"Workbook is unreadable or corrupt ({exc.__class__.__name__}): {detail}"],
        "warnings": [],
    }


def strict_json(value: str) -> Any:
    def pairs(items):
        result = {}
        for key, item in items:
            if key in result:
                raise ValueError(f"Duplicate JSON key: {key}")
            result[key] = item
        return result
    def invalid(value):
        raise ValueError(f"Non-finite JSON constant: {value}")
    return json.loads(value, object_pairs_hook=pairs, parse_constant=invalid)


def add_import_contract_errors(path, wb, data, errors, template_path=None, schema=None):
    schema = schema or EXPECTED_HEADERS
    from pipeline_ooxml import NS, Q, sheet_paths
    def string_values(parts):
        return [] if SHARED_STRINGS_PATH not in parts else [
            "".join(t.text or "" for t in si.iter(Q("t")))
            for si in ET.fromstring(parts[SHARED_STRINGS_PATH])]
    def semantic_tree(node, pool):
        if node is None:
            return None
        attrs = dict(node.attrib)
        children = list(node)
        if node.tag == Q("c") and attrs.get("t") in {"s", "inlineStr"}:
            value = node.find(Q("v"))
            if attrs["t"] == "inlineStr":
                resolved = "".join(t.text or "" for t in node.iter(Q("t")))
                attrs["t"] = "s"
            else:
                resolved = pool[int(value.text)] if value is not None and value.text else None
            return node.tag, tuple(sorted(attrs.items())), resolved
        text = node.text if node.tag in {Q("t"), Q("v")} else (node.text or "").strip()
        return node.tag, tuple(sorted(attrs.items())), text, tuple(semantic_tree(c, pool) for c in children)
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    paths = sheet_paths(entries)
    strings = string_values(entries)
    original = None
    if template_path:
        with zipfile.ZipFile(template_path) as archive:
            original = {name: archive.read(name) for name in archive.namelist()}
        original_paths = sheet_paths(original)
        old_strings = string_values(original)
        from build_pipeline_excel import ensure_shared_strings_content_type, ensure_shared_strings_relationship, ROW_KEY_HEADERS
        if Path(path).resolve() == Path(template_path).resolve():
            errors.append("Original template must be separate from the candidate workbook")
        allowed_new = {SHARED_STRINGS_PATH} if SHARED_STRINGS_PATH not in original else set()
        for name in entries.keys() - original.keys() - allowed_new:
            errors.append(f"Unexpected template package part added: {name}")
        for name, content in original.items():
            mutable = set(original_paths.values()) | {SHARED_STRINGS_PATH}
            if SHARED_STRINGS_PATH not in original:
                if name == "[Content_Types].xml":
                    content = ensure_shared_strings_content_type(content)
                elif name == "xl/_rels/workbook.xml.rels":
                    content = ensure_shared_strings_relationship(content)
            if name not in mutable and entries.get(name) != content:
                errors.append(f"Template package part changed: {name}")
        for name in paths.keys() - EXPECTED_HEADERS.keys():
            # Exports use shared strings. Synthetic inline templates require conversion.
            source = original.get(original_paths.get(name, ""), b"")
            if semantic_tree(ET.fromstring(entries[paths[name]]), strings) != semantic_tree(ET.fromstring(source), old_strings):
                errors.append(f"Template non-data sheet changed: {name}")
    for sheet, expected in schema.items():
        if sheet not in paths or sheet not in wb:
            continue
        root = ET.fromstring(entries[paths[sheet]])
        rows = root.find(Q("sheetData"))
        if rows is None:
            errors.append(f"{sheet}: missing sheetData")
            continue
        style_rows, keyed_style_rows, key_columns = {}, {}, []
        if original and sheet in original_paths:
            old_root = ET.fromstring(original[original_paths[sheet]])
            old_rows = old_root.find(Q("sheetData"))
            old_last = max((int(r.get("r")) for r in old_rows), default=2)
            from pipeline_ooxml import column_letter
            key_columns = [column_letter(expected.index(header) + 1) for header in ROW_KEY_HEADERS[sheet]]
            style_rows = {int(row.get("r")): row for row in old_rows}
            for number, old_row in style_rows.items():
                if number < 3:
                    continue
                old_cells = {''.join(c for c in cell.get('r') if c.isalpha()): cell for cell in old_row}
                key = tuple(semantic_tree(old_cells[column], old_strings)[2]
                            if column in old_cells and old_cells[column].get('t') in {'s', 'inlineStr'} else ''
                            for column in key_columns)
                if any(key):
                    keyed_style_rows[key] = old_row
            new_last = max((int(r.get("r")) for r in rows), default=2)
            if new_last < old_last:
                errors.append(f"{sheet}: template physical row boundary shrank ({old_last} -> {new_last})")
            old_shell, new_shell = deepcopy(old_root), deepcopy(root)
            old_shell.remove(old_shell.find(Q("sheetData")))
            new_shell.remove(new_shell.find(Q("sheetData")))
            if new_last > old_last:
                # Growth may extend dimension; columns, merges, panes and all other
                # worksheet layout remain the original export's contract.
                from pipeline_ooxml import column_letter
                new_dimension = new_shell.find(Q("dimension"))
                old_dimension = old_shell.find(Q("dimension"))
                if new_dimension is not None and old_dimension is not None:
                    if new_dimension.get("ref") != f"A1:{column_letter(len(expected))}{new_last}":
                        errors.append(f"{sheet}: invalid dimension extension")
                    new_dimension.attrib.clear()
                    new_dimension.attrib.update(old_dimension.attrib)
            if semantic_tree(old_shell, old_strings) != semantic_tree(new_shell, strings):
                errors.append(f"{sheet}: template worksheet layout changed (columns/merges/panes/dimensions)")
            for number in (1, 2):
                old_row = next((r for r in old_rows if r.get("r") == str(number)), None)
                new_row = next((r for r in rows if r.get("r") == str(number)), None)
                if old_row is not None and semantic_tree(old_row, old_strings) != semantic_tree(new_row, strings):
                    errors.append(f"{sheet}: template header/summary row {number} changed")
        for row in rows:
            number = int(row.get("r"))
            active = number >= 3 and any(clean(wb[sheet].cell(number, c).value) for c in range(1, len(expected) + 1))
            expected_styles = {}
            if number >= 3 and style_rows:
                key = tuple(str(wb[sheet][f'{column}{number}'].value or '') for column in key_columns)
                exemplar = keyed_style_rows.get(key) if active else None
                if exemplar is None:
                    exemplar = style_rows.get(number, style_rows.get(3))
                if exemplar is not None:
                    expected_styles = {''.join(c for c in cell.get('r') if c.isalpha()): cell.get('s') for cell in exemplar}
                    expected_attributes = dict(exemplar.attrib, r=str(number))
                    if row.attrib != expected_attributes:
                        errors.append(f"{sheet} row {number}: template row format changed")
            for cell in row:
                coord = cell.get("r", "?")
                column_name = ''.join(c for c in coord if c.isalpha())
                if number >= 3 and style_rows and cell.get('s') != expected_styles.get(column_name):
                    errors.append(f"{sheet} {coord}: template business-cell style changed")
                kind = cell.get("t")
                value = cell.find(Q("v"))
                if cell.find(Q("f")) is not None and number >= 3:
                    errors.append(f"{sheet} {coord}: formula is not a literal import value")
                if kind == "s":
                    try:
                        i = int(value.text) if value is not None else -1
                        if not 0 <= i < len(strings):
                            raise ValueError()
                    except (ValueError, TypeError):
                        errors.append(f"{sheet} {coord}: missing/invalid shared-string index")
                if number >= 3 and not active and (kind is not None or len(cell)):
                    errors.append(f"{sheet} {coord}: unused row must contain untyped true blanks")
                if number >= 3 and active and isinstance(wb[sheet][coord].value, str) and wb[sheet][coord].value and kind != "s":
                    errors.append(f"{sheet} {coord}: populated text must use shared strings")
                if number >= 3 and active and coord.startswith("A") and clean(wb[sheet][coord].value) == "":
                    errors.append(f"{sheet} {coord}: empty record key")
            if number >= 3:
                for column, header in enumerate(expected, 1):
                    value = wb[sheet].cell(number, column).value
                    if isinstance(value, str) and value and value != value.strip() and not header.endswith("_mlp_params"):
                        errors.append(f"{sheet} row {number}: leading/trailing whitespace in {header}")
                    if isinstance(value, str) and value.strip().lower() in {"<todo>", "<tbd>", "tbd", "todo", "待填写", "待确认"}:
                        errors.append(f"{sheet} row {number}: placeholder in {header}")


def validate_workbook(path: Path, template_path: Path | None = None,
                      business_reference_path: Path | None = None) -> dict[str, Any]:
    try:
        wb = load_workbook(path, data_only=False)
    except (OSError, KeyError, IndexError, TypeError, ValueError, ET.ParseError, zipfile.BadZipFile, InvalidFileException) as exc:
        return unreadable_workbook_result(path, exc)
    errors: list[str] = []
    warnings: list[str] = []
    row_counts: dict[str, int] = {}
    data: dict[str, list[dict[str, str]]] = {}

    add_ooxml_storage_errors(path, errors)
    from java_import_contract import template_headers, validate_java_contract, SOURCE_SHA256
    try:
        schema = template_headers(wb, EXPECTED_HEADERS)
    except ValueError as exc:
        errors.append(str(exc))
        schema = EXPECTED_HEADERS

    for sheet, expected in schema.items():
        if sheet not in wb.sheetnames:
            errors.append(f"Missing sheet: {sheet}")
            continue
        ws = wb[sheet]
        if any(ws.cell(2, c).value not in (None, "") for c in range(len(expected) + 1, ws.max_column + 1)):
            errors.append(f"{sheet}: unsupported additional header columns")
        actual = headers(ws, len(expected))
        if actual != expected:
            errors.append(f"{sheet} headers mismatch: expected={expected!r} actual={actual!r}")
            continue
        rows = row_dicts(ws, expected)
        data[sheet] = rows
        row_counts[sheet] = len(rows)
        add_empty_inline_string_errors(ws, len(expected), errors)
        for row in rows:
            for header in required_headers(expected):
                if not row.get(header):
                    errors.append(f"{sheet} row {row['_row']} missing required column {header}")

    dus = data.get(DATA_UTILIZATION_SHEET, [])
    targets = data.get(TARGET_SHEET, [])
    fields = data.get(FIELD_SHEET, [])
    pipelines = data.get(PIPELINE_SHEET, [])

    du_names = {row["*data_utilization_name"] for row in dus if row.get("*data_utilization_name")}
    target_names = {row["*target_name"] for row in targets if row.get("*target_name")}
    default_du_name = sorted(du_names)[0] if du_names else ""

    add_duplicate_errors([row.get("*data_utilization_name", "") for row in dus], "data_utilization_name", errors)
    add_duplicate_errors([row.get("*target_name", "") for row in targets], "target_name", errors)
    add_duplicate_errors([row.get("*data_utilization_name", "") + "/" + row.get("*pipeline_name", "") for row in pipelines], "data_utilization/pipeline_name", errors)
    add_duplicate_errors([row.get("dataset_name", "") for row in targets], "dataset_name", errors)
    add_duplicate_errors([row.get("*target_name", "") + "/" + row.get("*field_name", "") for row in fields], "target/field key", errors)
    sequences = defaultdict(int)
    for row in fields:
        target = row.get("*target_name")
        if row.get("field_sequence") != str(sequences[target]):
            errors.append(f"Target Field row {row['_row']}: field_sequence must be consecutive zero-based text per target")
        sequences[target] += 1

    for row in targets:
        if row.get("*data_utilization_name") not in du_names:
            errors.append(f"Target & Catalog row {row['_row']} references unknown data_utilization_name: {row.get('*data_utilization_name')}")
        storage = row.get("data_storage_type", "")
        if not row.get("dataset_name") and any(row.get(header) for header in EXPECTED_HEADERS[TARGET_SHEET][8:]):
            errors.append(f"Target & Catalog row {row['_row']}: Catalog metadata requires dataset_name")
        if storage not in ALLOWED_STORAGE_TYPES:
            errors.append(f"Target & Catalog row {row['_row']} has unsupported data_storage_type: {storage}")
        if row.get("dataset_name") and not row.get("dataset_business_owner_email"):
            warnings.append(f"Target & Catalog row {row['_row']} dataset owner/email fields are blank.")
        for name_column in [
            "dataset_business_owner",
            "dataset_it_owner",
            "dataset_fe",
            "dataset_it_bp",
            "dataset_data_engineer",
        ]:
            name_value = row.get(name_column, "")
            if "@" in name_value or any(ch.isdigit() for ch in name_value):
                warnings.append(f"Target & Catalog row {row['_row']} {name_column} should be a person-style name without email or digits.")

    for row in fields:
        target = row.get("*target_name")
        if target not in target_names:
            errors.append(f"Target Field row {row['_row']} references unknown target_name: {target}")
        if row.get("*field_label") != row.get("*field_name"):
            errors.append(
                f"Target Field row {row['_row']} field_label must equal field_name, "
                f"got field_name={row.get('*field_name')} field_label={row.get('*field_label')}"
            )
        raw_field_type = row.get("*field_type")
        if clean(raw_field_type) != REQUIRED_FIELD_TYPE:
            errors.append(f"Target Field row {row['_row']} field_type must be TEXT, got: {raw_field_type}")
        raw_field_length = row.get("field_length")
        if clean(raw_field_length) != REQUIRED_FIELD_LENGTH:
            errors.append(f"Target Field row {row['_row']} field_length must be 200, got: {raw_field_length}")

    if FIELD_SHEET in wb.sheetnames:
        field_ws = wb[FIELD_SHEET]
        length_col = EXPECTED_HEADERS[FIELD_SHEET].index("field_length") + 1
        sequence_col = EXPECTED_HEADERS[FIELD_SHEET].index("field_sequence") + 1
        for row_idx in range(3, field_ws.max_row + 1):
            length_cell = field_ws.cell(row_idx, length_col)
            if length_cell.value not in (None, "") and length_cell.data_type != "s":
                errors.append(f"Target Field row {row_idx} field_length must be stored as text, got cell type: {length_cell.data_type}")
            cell = field_ws.cell(row_idx, sequence_col)
            if cell.value in (None, ""):
                continue
            if cell.data_type != "s":
                errors.append(f"Target Field row {row_idx} field_sequence must be stored as text, got cell type: {cell.data_type}")

    target_field_counts: dict[str, int] = {}
    for row in fields:
        target = row.get("*target_name", "")
        target_field_counts[target] = target_field_counts.get(target, 0) + 1
    for target in sorted(target_names):
        target_row = next((row for row in targets if row.get("*target_name") == target), {})
        if target_field_counts.get(target, 0) == 0 and is_project_owned_table(target_row.get("table_name", ""), target_row.get("*data_utilization_name", "") or default_du_name):
            warnings.append(f"Target has no field rows: {target}")

    for row in pipelines:
        if row.get("*data_utilization_name") not in du_names:
            errors.append(f"Pipeline row {row['_row']} references unknown data_utilization_name: {row.get('*data_utilization_name')}")
        if row.get("*enable") not in {"0", "1"}:
            errors.append(f"Pipeline row {row['_row']} has unsupported enable value: {row.get('*enable')}")
        if row.get("*is_octopus") not in {"0", "1"}:
            errors.append(f"Pipeline row {row['_row']} has unsupported is_octopus value: {row.get('*is_octopus')}")
        if not row.get("task1_link_target_names"):
            warnings.append(f"Pipeline row {row['_row']} task1_link_target_names is blank.")
        if not row.get("task1_mlp_params"):
            warnings.append(f"Pipeline row {row['_row']} task1_mlp_params is blank.")

    try:
        add_import_contract_errors(path, wb, data, errors, template_path, schema)
        validate_java_contract(path, wb, data, schema, errors, warnings, strict_json)
        if business_reference_path is not None:
            from pipeline_preservation import preservation_errors
            if Path(path).resolve() == Path(business_reference_path).resolve():
                errors.append("Business reference must be the separate pre-edit workbook, not the output itself")
            else:
                errors.extend(preservation_errors(path, business_reference_path))
    except (OSError, KeyError, IndexError, TypeError, ValueError, ET.ParseError, zipfile.BadZipFile) as exc:
        errors.append(f"OOXML import preflight could not complete: {exc}")
    wb.close()
    return {
        "source_sha256": SOURCE_SHA256,
        "template_comparison": {"performed": template_path is not None,
                                "sha256": hashlib.sha256(Path(template_path).read_bytes()).hexdigest() if template_path and Path(template_path).is_file() else None},
        "business_preservation_checked": business_reference_path is not None,
        "limitations": ["XmlParameterHandler schema, enum definitions, datastore UID mapping and service/database side effects were not executed."],
        "ok": not errors,
        "workbook": str(path),
        "row_counts": row_counts,
        "errors": errors,
        "warnings": warnings,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Validate a Pipeline Export Excel workbook.")
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--json-out", type=Path)
    parser.add_argument("--template-xlsx", type=Path, help="Original export for row-boundary and package comparison")
    parser.add_argument("--business-reference-xlsx", type=Path, help="Pre-edit workbook for exact value/type/blank/XML preservation during format-only repairs")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    result = validate_workbook(args.xlsx, template_path=args.template_xlsx,
                               business_reference_path=args.business_reference_xlsx)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
