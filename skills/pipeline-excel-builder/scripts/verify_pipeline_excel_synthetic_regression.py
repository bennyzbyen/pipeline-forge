#!/usr/bin/env python3
"""Exercise Pipeline Excel in-place/copy writes against a synthetic OOXML template."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import tempfile
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

from build_pipeline_excel import EXPECTED_HEADERS, build_workbook
from validate_pipeline_excel import validate_workbook


SHEET_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def create_template(path: Path) -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Readme"
    readme["A1"] = "Synthetic Pipeline Export template"
    for sheet_name, headers in EXPECTED_HEADERS.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.cell(1, 1).value = "Synthetic fixture only"
        for column, header in enumerate(headers, start=1):
            worksheet.cell(2, column).value = header
        if sheet_name == "Pipeline":
            for column in (6, 7, 8):
                worksheet.cell(3, column).value = ""
    workbook.save(path)


def synthetic_facts() -> dict[str, object]:
    return {
        "data_utilizations": [
            {"data_utilization": "demo_quality", "description": "Synthetic quality output"}
        ],
        "report_targets": [
            {
                "data_utilization": "demo_quality",
                "target_name": "clickhouse_demo_result",
                "physical_table": "demo.demo_result",
                "storage": "ClickHouse",
                "description": "Synthetic target",
            }
        ],
        "report_physical_targets": [
            {
                "table": "demo.demo_result",
                "storage": "ClickHouse",
                "description": "Synthetic target",
            }
        ],
        "report_field_mappings": [
            {
                "inferred_target_name": "clickhouse_demo_result",
                "source_doc_index": 0,
                "fields": [
                    {"target_field": "record_id", "field_name": "Synthetic record identifier"},
                    {"target_field": "status_text", "field_name": "Synthetic status"},
                ],
            }
        ],
        "report_schedules": [
            {
                "data_utilization": "demo_quality",
                "pipeline_name": "daily_demo_result",
                "task_name": "build_demo_result",
                "description": "Synthetic daily pipeline",
            }
        ],
    }


def build_args(template: Path, facts: Path, questions: Path, output: Path | None) -> SimpleNamespace:
    return SimpleNamespace(
        template_xlsx=template,
        structured_facts=facts,
        out_xlsx=output,
        questions_out=questions,
        project_name="Synthetic Pipeline",
        data_utilization="",
    )


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def assert_ooxml_contract(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path, "r") as archive:
        names = set(archive.namelist())
        assert "xl/sharedStrings.xml" in names, names
        shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        shared_values = [
            "".join(element.itertext())
            for element in shared_root.findall(f"{{{SHEET_MAIN_NS}}}si")
        ]
        for required in ("TEXT", "200", "0", "1"):
            assert required in shared_values, (required, shared_values)

        worksheet_names = sorted(
            name for name in names if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
        for worksheet_name in worksheet_names:
            xml = archive.read(worksheet_name).decode("utf-8")
            assert 't="inlineStr"' not in xml, worksheet_name

    workbook = load_workbook(path, data_only=False)
    field_sheet = workbook["Target Field"]
    expected_field_cells = {
        "E3": "TEXT",
        "F3": "200",
        "G3": "0",
        "E4": "TEXT",
        "F4": "200",
        "G4": "1",
    }
    for coordinate, expected in expected_field_cells.items():
        cell = field_sheet[coordinate]
        assert cell.value == expected, (coordinate, cell.value)
        assert cell.data_type == "s", (coordinate, cell.data_type)

    pipeline_sheet = workbook["Pipeline"]
    for coordinate in ("F3", "G3", "H3"):
        cell = pipeline_sheet[coordinate]
        assert cell.value is None, (coordinate, cell.value)
        assert cell.data_type != "inlineStr", (coordinate, cell.data_type)

    return {
        "shared_string_count": len(shared_values),
        "inline_string_count": 0,
        "blank_timing_cells": 3,
        "fixed_text_cells": len(expected_field_cells),
    }


def assert_result(path: Path, expected_counts: dict[str, int]) -> dict[str, object]:
    validation = validate_workbook(path)
    assert validation["ok"], validation
    assert validation["row_counts"] == expected_counts, validation
    ooxml = assert_ooxml_contract(path)
    return {
        "status": "passed",
        "row_counts": validation["row_counts"],
        "warning_count": len(validation["warnings"]),
        **ooxml,
    }


def run_regression() -> dict[str, object]:
    expected_counts = {
        "Data Utilization": 1,
        "Target & Catalog": 1,
        "Target Field": 2,
        "Pipeline": 1,
    }
    with tempfile.TemporaryDirectory(prefix="pipeline_excel_synthetic_") as temp_dir:
        root = Path(temp_dir)
        facts_path = root / "structured_facts.json"
        facts_path.write_text(
            json.dumps(synthetic_facts(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        in_place_path = root / "in_place.xlsx"
        create_template(in_place_path)
        before_in_place = file_hash(in_place_path)
        counts = build_workbook(
            build_args(in_place_path, facts_path, root / "in_place_questions.md", None)
        )
        assert counts == expected_counts, counts
        assert file_hash(in_place_path) != before_in_place, "in-place mode did not update its template"
        in_place_result = assert_result(in_place_path, expected_counts)

        copy_template_path = root / "copy_template.xlsx"
        shutil.copyfile(in_place_path, copy_template_path)
        copy_template_hash = file_hash(copy_template_path)
        copy_output_path = root / "copy_output.xlsx"
        counts = build_workbook(
            build_args(
                copy_template_path,
                facts_path,
                root / "copy_questions.md",
                copy_output_path,
            )
        )
        assert counts == expected_counts, counts
        assert file_hash(copy_template_path) == copy_template_hash, "copy mode changed its source template"
        assert copy_output_path.is_file(), "copy mode did not create its requested output"
        copy_result = assert_result(copy_output_path, expected_counts)

    return {
        "status": "passed",
        "case_count": 2,
        "cases": [
            {"case": "in_place", **in_place_result},
            {"case": "copy", **copy_result},
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    print(json.dumps(run_regression(), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
