#!/usr/bin/env python3
"""Regression-test Pipeline Excel CLI failures against malformed local inputs."""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

from build_pipeline_excel import EXPECTED_HEADERS


SCRIPT_ROOT = Path(__file__).resolve().parent
VALIDATOR = SCRIPT_ROOT / "validate_pipeline_excel.py"
BUILDER = SCRIPT_ROOT / "build_pipeline_excel.py"


def run_cli(arguments: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, *arguments],
        text=True,
        encoding="utf-8",
        capture_output=True,
        check=False,
    )


def create_template(path: Path, *, missing_sheet: str = "", drift_sheet: str = "") -> None:
    workbook = Workbook()
    workbook.active.title = "Readme"
    for sheet_name, headers in EXPECTED_HEADERS.items():
        if sheet_name == missing_sheet:
            continue
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.cell(1, 1).value = "Synthetic malformed-input fixture"
        for column, header in enumerate(headers, start=1):
            worksheet.cell(2, column).value = header
        if sheet_name == drift_sheet:
            worksheet.cell(2, 1).value = "drifted_header"
    workbook.save(path)


def assert_failed_without_traceback(completed: subprocess.CompletedProcess[str]) -> str:
    combined = completed.stdout + completed.stderr
    assert completed.returncode != 0, combined
    assert "Traceback (most recent call last)" not in combined, combined
    assert combined.strip(), "failure was silent"
    return combined


def run_regression() -> dict[str, object]:
    results: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="pipeline_excel_input_robustness_") as temp_dir:
        root = Path(temp_dir)

        corrupt = root / "corrupt.xlsx"
        corrupt.write_bytes(b"this is not an OOXML zip archive")
        corrupt_report = root / "corrupt_validation.json"
        completed = run_cli([str(VALIDATOR), "--xlsx", str(corrupt), "--json-out", str(corrupt_report)])
        combined = assert_failed_without_traceback(completed)
        report = json.loads(corrupt_report.read_text(encoding="utf-8"))
        assert report["ok"] is False, report
        assert any("unreadable or corrupt" in error for error in report["errors"]), report
        assert "BadZipFile" in combined, combined
        results.append({"case": "corrupt_zip", "status": "passed", "exit_code": completed.returncode})

        missing = root / "missing_sheet.xlsx"
        create_template(missing, missing_sheet="Target Field")
        missing_report = root / "missing_sheet_validation.json"
        completed = run_cli([str(VALIDATOR), "--xlsx", str(missing), "--json-out", str(missing_report)])
        assert_failed_without_traceback(completed)
        report = json.loads(missing_report.read_text(encoding="utf-8"))
        assert "Missing sheet: Target Field" in report["errors"], report
        results.append({"case": "missing_required_sheet", "status": "passed", "exit_code": completed.returncode})

        drifted = root / "drifted_template.xlsx"
        create_template(drifted, drift_sheet="Target Field")
        original_hash = hashlib.sha256(drifted.read_bytes()).hexdigest()
        facts = root / "facts.json"
        facts.write_text("{}\n", encoding="utf-8")
        generated = root / "must_not_exist.xlsx"
        questions = root / "must_not_exist.md"
        completed = run_cli(
            [
                str(BUILDER),
                "--template-xlsx",
                str(drifted),
                "--structured-facts",
                str(facts),
                "--out-xlsx",
                str(generated),
                "--questions-out",
                str(questions),
                "--project-name",
                "Synthetic robustness fixture",
            ]
        )
        combined = assert_failed_without_traceback(completed)
        assert "headers mismatch" in combined, combined
        assert not generated.exists(), generated
        assert not questions.exists(), questions
        assert hashlib.sha256(drifted.read_bytes()).hexdigest() == original_hash, "builder changed drifted source"
        results.append({"case": "template_header_drift", "status": "passed", "exit_code": completed.returncode})

    return {
        "status": "passed",
        "case_count": len(results),
        "cases": results,
        "external_service_calls": 0,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    print(json.dumps(run_regression(), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
