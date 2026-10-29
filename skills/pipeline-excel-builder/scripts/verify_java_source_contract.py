"""Synthetic failure injections for the supplied Java handler's observable contracts."""
from copy import deepcopy
import unittest
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from build_pipeline_excel import rewrite_inline_strings_as_shared_strings
from pipeline_ooxml import Q, sheet_paths
from validate_pipeline_excel import validate_workbook
from verify_java_import_regression import ImportRegression


class JavaSourceContract(unittest.TestCase):
    setUp = ImportRegression.setUp
    build = ImportRegression.build
    mutate = ImportRegression.mutate
    errors = ImportRegression.errors

    def edit(self, sheet, values):
        with zipfile.ZipFile(self.output) as archive:
            infos = archive.infolist()
            parts = {info.filename: archive.read(info.filename) for info in infos}
        path = sheet_paths(parts)[sheet]
        root = ET.fromstring(parts[path])
        sst = ET.fromstring(parts["xl/sharedStrings.xml"])
        for coordinate, value in values.items():
            number = "".join(c for c in coordinate if c.isdigit())
            row = root.find(f"{Q('sheetData')}/{Q('row')}[@r='{number}']")
            cell = row.find(f"{Q('c')}[@r='{coordinate}']")
            if cell is None:
                cell = ET.SubElement(row, Q("c"), r=coordinate)
            for child in list(cell):
                cell.remove(child)
            cell.attrib.pop("t", None)
            if value is not None:
                cell.set("t", "s")
                ET.SubElement(cell, Q("v")).text = str(len(sst))
                ET.SubElement(ET.SubElement(sst, Q("si")), Q("t")).text = value
        sst.set("uniqueCount", str(len(sst)))
        parts[path] = ET.tostring(root, encoding="utf-8")
        parts["xl/sharedStrings.xml"] = ET.tostring(sst, encoding="utf-8")
        with zipfile.ZipFile(self.output, "w") as archive:
            for info in infos:
                archive.writestr(info, parts[info.filename])

    def test_final_record_without_sentinel_rejected(self):
        self.build()
        def truncate(root):
            data = root.find(Q("sheetData"))
            for row in list(data):
                if int(row.get("r")) > 3:
                    data.remove(row)
        self.mutate("Data Utilization", truncate)
        self.assertIn("blank sentinel", self.errors())

    def test_missing_physical_row_would_null_dereference(self):
        self.build()
        def remove(root):
            data = root.find(Q("sheetData"))
            data.remove(data.find(f"{Q('row')}[@r='8']"))
        self.mutate("Pipeline", remove)
        self.assertIn("isRowEmpty(null)", self.errors())

    def test_sheet_order_is_positional(self):
        wb = load_workbook(self.template)
        wb.move_sheet("Pipeline", offset=-3)
        wb.save(self.template)
        wb.close()
        with self.assertRaisesRegex(ValueError, "sheet position"):
            self.build()

    def test_empty_string_conditional_is_not_java_null(self):
        self.build()
        self.edit("Target & Catalog", {"H3": ""})
        self.edit("Pipeline", {"F3": ""})
        self.assertGreaterEqual(self.errors().count("empty STRING is not null"), 2)

    def test_cross_du_link_rejected(self):
        self.facts["report_targets"].append({"data_utilization": "second", "target_name": "foreign", "storage": "HBASE", "physical_table": "foreign"})
        self.facts["report_schedules"][0]["task1_link_target_names"] = ["foreign"]
        with self.assertRaisesRegex(ValueError, "cross-data_utilization"):
            self.build()

    def test_same_pipeline_name_allowed_in_different_du(self):
        self.facts["report_targets"].append({"data_utilization": "second", "target_name": "foreign", "storage": "HBASE", "physical_table": "foreign"})
        other = deepcopy(self.facts["report_schedules"][0])
        other.update(data_utilization="second", task1_link_target_names=["foreign"])
        self.facts["report_schedules"].append(other)
        self.build(in_place=True)
        self.build()
        self.assertEqual(validate_workbook(self.output)["row_counts"]["Pipeline"], 2)

    def test_storage_conditional_requirements(self):
        self.build()
        self.edit("Target & Catalog", {"E3": None})
        self.assertIn("missing db", self.errors())
        self.edit("Target & Catalog", {"D3": "HBASE"})
        self.assertNotIn("missing db", self.errors())
        self.edit("Target & Catalog", {"F3": None})
        self.assertIn("missing table_name", self.errors())

    def test_mysql_not_supported_by_persistence_switch(self):
        self.build()
        self.edit("Target & Catalog", {"D3": "MYSQL"})
        self.assertIn("no supported upsertDataStorage branch", self.errors())

    def test_dataset_requires_all_owner_and_description_fields(self):
        self.build()
        self.edit("Target & Catalog", {"H3": "catalog"})
        result = validate_workbook(self.output)
        conditional = [e for e in result["errors"] if "conditional requirement" in e]
        self.assertEqual(len(conditional), 12)

    def test_dataset_source_xml_checked(self):
        self.build()
        self.edit("Target & Catalog", {"U3": "<Datasources>"})
        self.assertIn("dataset_source: invalid XML/JSON", self.errors())

    def test_real_schedule_format_and_dates(self):
        for trigger in ("每天 03:00", "每2天 9:05"):
            with self.subTest(trigger=trigger):
                self.facts["report_schedules"][0].update(pipeline_trigger=trigger, pipeline_trigger_start="2027-01-01")
                self.build()
                self.assertTrue(validate_workbook(self.output)["ok"], self.errors())
        self.edit("Pipeline", {"F3": "0 0 3 * * ?"})
        self.assertIn("not cron", self.errors())

    def test_schedule_start_and_end_errors_collected(self):
        self.build()
        self.edit("Pipeline", {"F3": "每天 03:00", "G3": None, "H3": "2027-02-30"})
        self.assertIn("requires pipeline_trigger_start", self.errors())
        self.assertIn("valid yyyy-MM-dd", self.errors())
        self.edit("Pipeline", {"G3": "2027-03-01", "H3": "2027-02-28"})
        self.assertIn("precedes", self.errors())

    def test_dates_parsed_even_without_trigger(self):
        self.build()
        self.edit("Pipeline", {"G3": "not-a-date"})
        self.assertIn("valid yyyy-MM-dd", self.errors())

    def test_notification_email_and_enum_limits(self):
        self.build()
        self.edit("Pipeline", {"F3": "每天 03:00", "G3": "2027-01-01", "I3": "FAILURE, SUCCESS", "J3": "ops+tag@example.com"})
        self.assertIn("does not trim", self.errors())
        self.assertIn("email addresses", self.errors())
        self.edit("Pipeline", {"I3": "FAILURE", "J3": "ops@example.com"})
        result = validate_workbook(self.output)
        self.assertTrue(result["ok"], result["errors"])
        self.assertTrue(any("AsyncJobStatusEnum" in w for w in result["warnings"]))

    def test_fullwidth_link_comma_rejected(self):
        self.build()
        self.edit("Pipeline", {"M3": "clickhouse_demo_result，clickhouse_demo_result"})
        self.assertIn("split links differently", self.errors())

    def test_linked_task_missing_xml_and_orphan_task(self):
        self.build()
        self.edit("Pipeline", {"N3": None, "K3": None})
        self.assertIn("calls getMlpParameter", self.errors())
        self.assertIn("silently discard", self.errors())

    def test_two_tasks_preserved_and_second_task_validated(self):
        wb = load_workbook(self.template)
        for column, suffix in enumerate(("name", "description", "link_target_names", "mlp_params"), 15):
            wb["Pipeline"].cell(2, column, "task2_" + suffix)
        wb.save(self.template)
        wb.close()
        rewrite_inline_strings_as_shared_strings(self.template)
        task = {"name": "execute", "description": "Synthetic task", "link_target_names": ["clickhouse_demo_result"], "mlp_params": self.params}
        self.facts["report_schedules"][0]["tasks"] = [task, deepcopy(task)]
        self.build()
        wb = load_workbook(self.output)
        self.assertEqual(wb["Pipeline"]["R3"].value, self.params)
        wb.close()
        self.edit("Pipeline", {"R3": "<bad>"})
        self.assertIn("task2 MLP: invalid XML/JSON", self.errors())

    def test_task_capacity_error_preserves_original(self):
        before = self.template.read_bytes()
        self.facts["report_schedules"][0]["tasks"] = [{"name": "execute", "link_target_names": [], "mlp_params": self.params}] * 2
        with self.assertRaisesRegex(ValueError, "insufficient task column groups"):
            self.build(in_place=True)
        self.assertEqual(self.template.read_bytes(), before)

    def test_values_beyond_declared_headers_are_not_silently_lost(self):
        self.build()
        self.edit("Pipeline", {"O3": "hidden task without header"})
        self.assertIn("outside the declared header", self.errors())


if __name__ == "__main__":
    unittest.main(verbosity=2)
