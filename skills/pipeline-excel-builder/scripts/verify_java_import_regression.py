"""Sanitized offline regressions for observed Java-import compatibility failures.

These checks do not emulate or claim to execute the proprietary Java importer.
"""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from build_pipeline_excel import (EXPECTED_HEADERS, build_workbook,
                                  rewrite_inline_strings_as_shared_strings)
from pipeline_ooxml import Q, sheet_paths
from validate_pipeline_excel import validate_workbook
from verify_pipeline_excel_synthetic_regression import create_template, synthetic_facts, build_args, file_hash


class ImportRegression(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="pipeline_contract_")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.template = self.root / "template.xlsx"
        self.output = self.root / "final.xlsx"
        create_template(self.template)
        wb = load_workbook(self.template)
        for sheet in EXPECTED_HEADERS:
            ws = wb[sheet]
            ws.row_dimensions[20].height = 23
            for column in range(1, len(EXPECTED_HEADERS[sheet]) + 1):
                ws.cell(20, column, "old value").fill = PatternFill("solid", fgColor="EEEEEE")
        wb.save(self.template)
        wb.close()
        rewrite_inline_strings_as_shared_strings(self.template)
        self.facts = synthetic_facts()
        self.facts["report_schedules"][0]["task1_link_target_names"] = ["clickhouse_demo_result"]
        self.params = '<MLP><class>example.jobs.ExecuteReport</class><params><![CDATA[{"year":2027,"tables":["demo_result"]}]]></params></MLP>'
        self.facts["report_schedules"][0]["task1_mlp_params"] = self.params

    def build(self, facts=None, in_place=False):
        source = self.root / "facts.json"
        source.write_text(json.dumps(facts or self.facts), encoding="utf-8")
        build_workbook(build_args(self.template, source, self.root / "questions.md", None if in_place else self.output))
        return self.template if in_place else self.output

    def mutate(self, sheet, change):
        with zipfile.ZipFile(self.output) as archive:
            infos = archive.infolist()
            parts = {i.filename: archive.read(i.filename) for i in infos}
        path = sheet_paths(parts)[sheet]
        root = ET.fromstring(parts[path])
        change(root)
        parts[path] = ET.tostring(root, encoding="utf-8")
        with zipfile.ZipFile(self.output, "w") as archive:
            for info in infos:
                archive.writestr(info, parts[info.filename])

    def errors(self):
        return "\n".join(validate_workbook(self.output, self.template)["errors"])

    def test_preserve_package_rows_and_true_blank_tail(self):
        before = file_hash(self.template)
        self.build()
        self.assertEqual(before, file_hash(self.template))
        self.assertTrue(validate_workbook(self.output, self.template)["ok"], self.errors())
        with zipfile.ZipFile(self.template) as source, zipfile.ZipFile(self.output) as output:
            for part in ("xl/styles.xml", "xl/workbook.xml", "xl/worksheets/sheet1.xml", "docProps/core.xml"):
                self.assertEqual(source.read(part), output.read(part))
            parts = {p: output.read(p) for p in output.namelist()}
            for name, path in sheet_paths(parts).items():
                if name not in EXPECTED_HEADERS:
                    continue
                tail = ET.fromstring(parts[path]).find(f"{Q('sheetData')}/{Q('row')}[@r='20']")
                self.assertEqual(tail.get("ht"), "23")
                self.assertTrue(all(not len(c) and c.get("t") is None for c in tail))

    def test_xml_cdata_and_execution_class_preserved_exactly(self):
        self.build()
        wb = load_workbook(self.output)
        self.assertEqual(wb["Pipeline"]["N3"].value, self.params)
        self.assertIsNone(wb["Pipeline"]["I3"].value)
        wb.close()

    def test_multiple_utilizations_and_repeated_task_names(self):
        facts = deepcopy(self.facts)
        target = deepcopy(facts["report_targets"][0])
        target.update(data_utilization="other", target_name="other_target", physical_table="other.result")
        facts["report_targets"].append(target)
        pipe = deepcopy(facts["report_schedules"][0])
        pipe.update(data_utilization="other", pipeline_name="unrelated_name", task1_link_target_names=["other_target"])
        facts["report_schedules"].append(pipe)
        self.build(facts)
        wb = load_workbook(self.output)
        self.assertEqual(wb["Target & Catalog"]["A4"].value, "other")
        self.assertEqual(wb["Pipeline"]["A4"].value, "other")
        self.assertTrue(validate_workbook(self.output)["ok"], self.errors())
        wb.close()

    def test_duplicate_dataset_requires_confirmed_binding(self):
        facts = deepcopy(self.facts)
        other = deepcopy(facts["report_targets"][0])
        other["target_name"] = "second_storage"
        facts["report_targets"].append(other)
        facts["report_catalog_basic_info"] = [{"data_item": "demo_result", "title": "Result", **{role: "Demo Owner <owner@example.com>" for role in ("business_owner", "it_owner", "fe", "it_bp", "data_engineer")}}]
        before = file_hash(self.template)
        with self.assertRaisesRegex(ValueError, "Duplicate dataset_name"):
            self.build(facts, in_place=True)
        self.assertEqual(before, file_hash(self.template))
        facts["report_catalog_basic_info"][0]["target_name"] = "second_storage"
        self.build(facts)
        self.assertTrue(validate_workbook(self.output)["ok"], self.errors())

    def test_no_schedule_drops_inherited_notification(self):
        self.facts["report_schedules"][0]["pipeline_status_notification"] = "SUCCESS,FAILURE"
        self.build()
        wb = load_workbook(self.output)
        self.assertIsNone(wb["Pipeline"]["I3"].value)
        wb.close()

    def test_explicit_schedule_and_notification_retained(self):
        self.facts["report_schedules"][0].update(pipeline_trigger="每天 03:00", pipeline_trigger_start="2027-01-01", pipeline_status_notification="FAILURE", pipeline_notification_emails="ops@example.com")
        self.build()
        wb = load_workbook(self.output)
        self.assertEqual(wb["Pipeline"]["I3"].value, "FAILURE")
        wb.close()

    def test_malformed_xml_and_strict_json_leave_source_and_output_unchanged(self):
        self.output.write_bytes(b"existing final")
        before = file_hash(self.template)
        for params in ("<MLP>", '<MLP><p>{"a":1,"a":2}</p></MLP>', '<MLP><p>{"a":NaN}</p></MLP>'):
            with self.subTest(params=params):
                self.facts["report_schedules"][0]["task1_mlp_params"] = params
                with self.assertRaisesRegex(ValueError, "XML/JSON"):
                    self.build()
                self.assertEqual(before, file_hash(self.template))
                self.assertEqual(self.output.read_bytes(), b"existing final")

    def test_unknown_or_duplicate_explicit_links_fail(self):
        for links in (["unknown"], ["clickhouse_demo_result"] * 2, [" clickhouse_demo_result"]):
            self.facts["report_schedules"][0]["task1_link_target_names"] = links
            with self.assertRaises(ValueError):
                self.build()

    def test_duplicate_input_identifiers_are_not_silently_dropped(self):
        for key in ("report_targets", "report_schedules"):
            facts = deepcopy(self.facts)
            facts[key].append(deepcopy(facts[key][0]))
            with self.assertRaisesRegex(ValueError, "duplicate"):
                self.build(facts)
        facts = deepcopy(self.facts)
        facts["report_field_mappings"][0]["fields"] *= 2
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.build(facts)

    def test_typed_empty_tail_is_reported(self):
        self.build()
        self.mutate("Data Utilization", lambda root: root.find(f"{Q('sheetData')}/{Q('row')}[@r='20']/{Q('c')}").set("t", "s"))
        errors = self.errors()
        self.assertIn("shared-string index", errors)
        self.assertIn("untyped true blanks", errors)

    def test_shrunken_physical_rows_are_reported(self):
        self.build()
        def shrink(root):
            data = root.find(Q("sheetData"))
            for row in list(data):
                if int(row.get("r")) > 3:
                    data.remove(row)
        self.mutate("Data Utilization", shrink)
        self.assertIn("boundary shrank", self.errors())

    def test_unknown_du_is_reported(self):
        self.build()
        def change(root):
            a3 = root.find(f"{Q('sheetData')}/{Q('row')}[@r='3']/{Q('c')}[@r='A3']/{Q('v')}")
            b3 = root.find(f"{Q('sheetData')}/{Q('row')}[@r='3']/{Q('c')}[@r='B3']/{Q('v')}")
            a3.text = b3.text
        self.mutate("Pipeline", change)
        self.assertIn("unknown data_utilization_name", self.errors())

    def test_duplicate_field_and_sequence_reported(self):
        self.build()
        def duplicate(root):
            data = root.find(Q("sheetData"))
            row = deepcopy(data.find(f"{Q('row')}[@r='3']"))
            row.set("r", "4")
            for cell in row:
                cell.set("r", cell.get("r").replace("3", "4"))
            existing = data.find(f"{Q('row')}[@r='4']")
            data.remove(existing)
            data.insert(3, row)
        self.mutate("Target Field", duplicate)
        self.assertIn("Duplicate target/field key", self.errors())
        self.assertIn("zero-based", self.errors())

    def test_field_dictionary_order_survives_existing_template(self):
        self.build(in_place=True)
        self.facts["report_field_mappings"][0]["fields"].reverse()
        self.build()
        wb = load_workbook(self.output)
        self.assertEqual(wb["Target Field"]["B3"].value, "status_text")
        self.assertEqual(wb["Target Field"]["G3"].value, "0")
        wb.close()

    def test_header_whitespace_and_unknown_columns_rejected(self):
        for bad_header in ("*data_utilization_name ", "task2_name"):
            with self.subTest(header=bad_header):
                wb = load_workbook(self.template)
                if bad_header == "task2_name":
                    wb["Pipeline"]["O2"] = bad_header
                else:
                    wb["Data Utilization"]["A2"] = bad_header
                wb.save(self.template)
                wb.close()
                with self.assertRaisesRegex(ValueError, "headers mismatch|tasks cols error"):
                    self.build()
                wb = load_workbook(self.template)
                wb["Data Utilization"]["A2"] = "*data_utilization_name"
                wb["Pipeline"]["O2"] = None
                wb.save(self.template)
                wb.close()

    def test_mixed_string_normalization_preserves_existing_indices(self):
        wb = load_workbook(self.template)
        wb["Readme"]["B1"] = "First value"
        wb.save(self.template)
        wb.close()
        rewrite_inline_strings_as_shared_strings(self.template)
        with zipfile.ZipFile(self.template) as archive:
            infos = archive.infolist()
            parts = {info.filename: archive.read(info.filename) for info in infos}
        path = sheet_paths(parts)["Readme"]
        root = ET.fromstring(parts[path])
        row = root.find(f"{Q('sheetData')}/{Q('row')}")
        cell = ET.SubElement(row, Q("c"), r="C1", t="inlineStr")
        ET.SubElement(ET.SubElement(cell, Q("is")), Q("t")).text = "Second value"
        # Serialize without a prefix so this is a supported inline fixture shape.
        parts[path] = ET.tostring(root, encoding="utf-8")
        with zipfile.ZipFile(self.template, "w") as archive:
            for info in infos:
                archive.writestr(info, parts[info.filename])
        rewrite_inline_strings_as_shared_strings(self.template)
        wb = load_workbook(self.template)
        self.assertEqual(wb["Readme"]["B1"].value, "First value")
        self.assertEqual(wb["Readme"]["C1"].value, "Second value")
        wb.close()

    def test_template_tail_pattern_survives_data_growth(self):
        self.build(in_place=True)  # now has a confirmed true blank tail
        for i in range(25):
            self.facts["report_field_mappings"][0]["fields"].append({"target_field": f"extra_{i}"})
        self.build()
        wb = load_workbook(self.output)
        ws = wb["Target Field"]
        self.assertEqual(ws.max_row, 30)  # 2 header + 27 data + 1 retained blank tail
        self.assertIsNone(ws["A30"].value)
        wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
