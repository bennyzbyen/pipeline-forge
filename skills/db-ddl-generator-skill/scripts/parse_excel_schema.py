#!/usr/bin/env python3
"""Parse CSV/XLSX field lists into the standard schema contract."""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional


HEADER_ALIASES = {
    "name": ["字段", "字段名", "字段英文名", "英文字段", "L3 字段名", "column_name", "column", "name", "field", "field_name"],
    "comment": ["字段描述", "字段中文名", "中文名", "中文注释", "comment", "description", "desc", "label", "remarks", "备注"],
    "type": ["字段类型", "数据类型", "类型", "type", "data_type", "datatype", "db_type"],
    "length": ["长度", "length", "len", "size"],
    "precision": ["精度", "precision"],
    "scale": ["小数位", "scale"],
    "nullable": ["是否为空", "可为空", "nullable", "null", "is_nullable"],
    "not_null": ["非空", "not_null", "not null", "required", "mandatory"],
    "default": ["默认值", "default", "default_value"],
    "primary_key": ["主键", "primary_key", "pk", "is_pk"],
    "index": ["索引", "index", "key", "idx"],
    "remarks": ["备注", "remark", "remarks", "说明"],
}

TRUTHY = {"1", "true", "yes", "y", "是", "有", "√", "yes.", "pk", "primary", "主键"}
FALSY = {"0", "false", "no", "n", "否", "无", "×", "null", "nullable"}
NOT_NULL_VALUES = {"非空", "not null", "not_null", "required", "mandatory", "必填", "否", "no", "n", "false", "0"}
NULLABLE_VALUES = {"可空", "空", "nullable", "null", "是", "yes", "y", "true", "1"}


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def norm_header(value: Any) -> str:
    return re.sub(r"[\s_\-()/（）\\]+", "", cell_text(value).lower())


def alias_lookup() -> Dict[str, str]:
    lookup: Dict[str, str] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            lookup[norm_header(alias)] = canonical
    return lookup


def canonical_header(value: Any) -> Optional[str]:
    normalized = norm_header(value)
    if not normalized:
        return None
    lookup = alias_lookup()
    if normalized in lookup:
        return lookup[normalized]
    aliases = sorted(lookup.items(), key=lambda item: len(item[0]), reverse=True)
    for alias, canonical in aliases:
        if normalized.startswith(alias):
            return canonical
    return None


def truthy(value: Any) -> bool:
    return cell_text(value).lower() in TRUTHY


def falsy(value: Any) -> bool:
    return cell_text(value).lower() in FALSY


def parse_bool(value: Any) -> bool:
    text = cell_text(value).lower()
    if text in TRUTHY:
        return True
    if text in FALSY or not text:
        return False
    return bool(text)


def parse_int(value: Any) -> Optional[int]:
    text = cell_text(value)
    if not text:
        return None
    match = re.search(r"\d+", text)
    return int(match.group(0)) if match else None


def split_type(type_text: str) -> Dict[str, Optional[int]]:
    result: Dict[str, Optional[int]] = {"length": None, "precision": None, "scale": None}
    match = re.search(r"\((\d+)(?:\s*,\s*(\d+))?\)", type_text)
    if not match:
        return result
    first = int(match.group(1))
    second = int(match.group(2)) if match.group(2) else None
    if second is None:
        result["length"] = first
    else:
        result["precision"] = first
        result["scale"] = second
    return result


def detect_header_row(rows: List[List[Any]]) -> Optional[int]:
    lookup = alias_lookup()
    best_idx = None
    best_score = 0
    for idx, row in enumerate(rows[:30]):
        seen = {canonical_header(cell) for cell in row if norm_header(cell)}
        score = len([x for x in seen if x])
        if "name" in seen and ("type" in seen or "comment" in seen) and score > best_score:
            best_idx = idx
            best_score = score
    return best_idx


def normalize_schema_from_rows(
    rows: List[List[Any]],
    table_name: Optional[str] = None,
    source: Optional[str] = None,
) -> Dict[str, Any]:
    warnings: List[str] = []
    header_idx = detect_header_row(rows)
    if header_idx is None:
        return {
            "table_name": table_name,
            "columns": [],
            "warnings": ["no recognizable header row with a column-name field"],
            "source": source,
        }

    header = rows[header_idx]
    candidates: Dict[str, List[int]] = {}
    for idx, cell in enumerate(header):
        canonical = canonical_header(cell)
        if canonical:
            candidates.setdefault(canonical, []).append(idx)

    positions: Dict[int, str] = {}
    data_rows = rows[header_idx + 1 :]
    for canonical, indexes in candidates.items():
        best_idx = max(
            indexes,
            key=lambda col_idx: sum(1 for row in data_rows if col_idx < len(row) and cell_text(row[col_idx])),
        )
        positions[best_idx] = canonical

    columns: List[Dict[str, Any]] = []
    primary_key: List[str] = []
    indexes: List[Dict[str, Any]] = []
    seen_names = set()

    for row in rows[header_idx + 1 :]:
        values = {field: cell_text(row[idx]) if idx < len(row) else "" for idx, field in positions.items()}
        name = values.get("name", "").strip()
        if not name or name in seen_names:
            continue
        if re.fullmatch(r"[-_=]+", name):
            continue
        seen_names.add(name)

        source_type = values.get("type", "").strip()
        if not source_type:
            source_type = "string"
            warnings.append(f"{name}: missing type, defaulted to string")

        primary = parse_bool(values.get("primary_key", ""))
        nullable_text = values.get("nullable", "")
        not_null_text = values.get("not_null", "")
        if primary:
            nullable = False
        elif cell_text(not_null_text).lower() in NOT_NULL_VALUES or truthy(not_null_text):
            nullable = False
        elif cell_text(nullable_text).lower() in NULLABLE_VALUES:
            nullable = True
        elif cell_text(nullable_text).lower() in NOT_NULL_VALUES:
            nullable = False
        else:
            nullable = True
            warnings.append(f"{name}: missing nullable flag, defaulted to nullable")

        parsed_type = split_type(source_type)
        length = parse_int(values.get("length")) or parsed_type["length"]
        precision = parse_int(values.get("precision")) or parsed_type["precision"]
        scale = parse_int(values.get("scale")) or parsed_type["scale"]

        column = {
            "name": name,
            "source_type": source_type,
            "type": source_type,
            "length": length,
            "precision": precision,
            "scale": scale,
            "nullable": nullable,
            "default": values.get("default") or None,
            "primary_key": primary,
            "index": parse_bool(values.get("index", "")),
            "comment": values.get("comment") or "",
            "remarks": values.get("remarks") or "",
        }
        columns.append(column)
        if primary:
            primary_key.append(name)
        if column["index"]:
            indexes.append({"name": f"idx_{table_name or 'table'}_{name}", "columns": [name]})

    return {
        "table_name": table_name,
        "columns": columns,
        "primary_key": primary_key,
        "indexes": indexes,
        "warnings": warnings,
        "source": source,
    }


def read_csv_rows(path: Path) -> List[List[Any]]:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [row for row in csv.reader(handle)]
        except UnicodeError:
            continue
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        return [row for row in csv.reader(handle)]


def read_xlsx_sheets(path: Path, sheet: Optional[str] = None) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required to parse .xlsx files") from exc

    # Embedded workbooks in DOCX may have a stale dimension ref like A1.
    # read_only=True trusts that dimension, so use normal mode to scan real cells.
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet_names = [sheet] if sheet else workbook.sheetnames
    sheets: List[Dict[str, Any]] = []
    for sheet_name in sheet_names:
        worksheet = workbook[sheet_name]
        rows = [[cell_text(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
        sheets.append({"sheet_name": sheet_name, "rows": rows})
    workbook.close()
    return sheets


def load_workbook_schemas(path: Path, table_name: Optional[str] = None, sheet: Optional[str] = None) -> List[Dict[str, Any]]:
    suffix = path.suffix.lower()
    schemas: List[Dict[str, Any]] = []
    if suffix in {".csv", ".tsv"}:
        rows = read_csv_rows(path)
        schemas.append(normalize_schema_from_rows(rows, table_name=table_name or path.stem, source=str(path)))
    elif suffix in {".xlsx", ".xlsm"}:
        for sheet_data in read_xlsx_sheets(path, sheet=sheet):
            schema = normalize_schema_from_rows(
                sheet_data["rows"],
                table_name=table_name or sheet_data["sheet_name"],
                source=f"{path}#{sheet_data['sheet_name']}",
            )
            if schema.get("columns"):
                schemas.append(schema)
    else:
        raise SystemExit(f"unsupported file type: {path.suffix}")
    return schemas


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--table", help="Target table name")
    parser.add_argument("--sheet", help="Excel sheet name")
    parser.add_argument("--output", type=Path, help="Write JSON output to this path")
    parser.add_argument("--all-sheets", action="store_true", help="Keep all detected schemas instead of only the first one")
    args = parser.parse_args()

    schemas = load_workbook_schemas(args.input, table_name=args.table, sheet=args.sheet)
    if not schemas:
        result: Dict[str, Any] = {
            "table_name": args.table or args.input.stem,
            "columns": [],
            "warnings": ["no field definition table detected"],
        }
    elif args.all_sheets:
        result = {"tables": schemas, "warnings": []}
    else:
        result = schemas[0]

    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(output + "\n", encoding="utf-8")
    else:
        print(output)


if __name__ == "__main__":
    main()
